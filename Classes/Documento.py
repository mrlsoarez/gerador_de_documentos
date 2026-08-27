
from openpyxl import load_workbook

import shutil
import locale

from docx import Document
from docx2pdf import convert
from copy import deepcopy

from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import date
from docx.shared import Pt

from Services.Data import EncontrarData

class Documento: 
    
    def __init__(self, info_processo, info_secretarias, info):
        self.info_processo = info_processo 
        self.info_secretarias = info_secretarias
        self.info = info
    
    def get_mapeamento():
        pass 
    
    def copiar_arquivo(self, antigo, novo):
        shutil.copy(antigo, novo)
    
    def mudar_fonte(texto, name_font):
        texto.font.name = name_font
    
    def mudar_tamanho(texto, px): 
        texto.font.size = Pt(px)

    @staticmethod
    def criar_texto(paragrafo, texto, negrito = False, posicionamento = None, px = None, fonte = None):

        def deixar_negrito(run): 
           run.bold = True 

        def alinhar_texto(paragrafo, alinhado):
            if (alinhado == "Centro"): 
                paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif (alinhado == "Direita"):
                paragrafo.alignment = WD_ALIGN_PARAGRAPH.RIGHT 
            elif (alinhado == "ESQUERDA"):
                                paragrafo.alignment = WD_ALIGN_PARAGRAPH.LEFT

        run = paragrafo.add_run(str(texto))
        if (negrito): deixar_negrito(run)
        if (posicionamento != None): alinhar_texto(paragrafo, posicionamento)
        if (px != None): Documento.mudar_tamanho(run, px)
        if (fonte != None): Documento.mudar_fonte(run, fonte)

    @staticmethod
    def encontrar_data_de_hoje_em_extenso():  

        dia = EncontrarData("dia")
        mes = EncontrarData("mes", True)
        ano = EncontrarData("ano")
        
        return f"{dia} de {mes} de {ano}"

    @staticmethod
    def adicionar_linha_de_assinatura(paragrafo, assinador): 
        Documento.criar_texto(paragrafo, f"________________________________________\n{assinador}", negrito = True, posicionamento = "Centro", fonte = "Arial")
    
    @staticmethod
    def modificar_tabela(tabela, dados, fonte = "Cambria"): 

        for i in range(len(dados)):

            dado = dados[i]
                
            celula = tabela.cell(dado["celula"][0], dado["celula"][1])
            paragrafo = celula.paragraphs[0]  

            Documento.criar_texto(paragrafo, dado["conteudo"], dado["negrito"], fonte, px = 10)

        return tabela
    

                   

        

 
        
